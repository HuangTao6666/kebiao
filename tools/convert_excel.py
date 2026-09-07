#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把教务系统导出的周课表 Excel 转成应用内置数据 data.js。
用法:
    python convert_excel.py 课表.xlsx -o data.js
    python convert_excel.py 课表.xlsx -o data.js --semester-start 2026-09-07 --days 7
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)

DAY_NAMES = [
    ["星期一", "礼拜一", "周一", "Monday", "Mon"],
    ["星期二", "礼拜二", "周二", "Tuesday", "Tue"],
    ["星期三", "礼拜三", "周三", "Wednesday", "Wed"],
    ["星期四", "礼拜四", "周四", "Thursday", "Thu"],
    ["星期五", "礼拜五", "周五", "Friday", "Fri"],
    ["星期六", "礼拜六", "周六", "Saturday", "Sat"],
    ["星期日", "星期天", "礼拜日", "周日", "Sunday", "Sun"],
]

NUM_CHAR = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


def match_day(text):
    s = str(text or "").strip()
    if not s:
        return 0
    for d, tokens in enumerate(DAY_NAMES, 1):
        for t in tokens:
            if s == t or s.upper() == t.upper():
                return d
    if re.fullmatch(r"[1-7]", s):
        return int(s)
    if s in ("日", "天"):
        return 7
    if s in NUM_CHAR and NUM_CHAR[s] <= 7:
        return NUM_CHAR[s]
    return 0


def parse_ranges(raw):
    s = str(raw or "").strip()
    odd_even = None
    m = re.search(r"[（(]\s*([单双])\s*(?:周|週)?\s*[)）]", s) or re.search(r"([单双])\s*(?:周|週)?\s*$", s)
    if m:
        odd_even = "odd" if m.group(1) == "单" else "even"
        s = s.replace(m.group(0), "")
    s = re.sub(r"[周週节節]", "", s)
    values = []
    for part in re.split(r"[,，、;；\s]+", s):
        if not part:
            continue
        rm = re.fullmatch(r"(\d+)\s*[-–—~～]\s*(\d+)", part)
        if rm:
            a, b = int(rm.group(1)), int(rm.group(2))
            values.extend(range(min(a, b), max(a, b) + 1))
        elif re.fullmatch(r"\d+", part):
            values.append(int(part))
    values = sorted(set(values))
    if odd_even == "odd":
        values = [n for n in values if n % 2 == 1]
    elif odd_even == "even":
        values = [n for n in values if n % 2 == 0]
    return values


def summarize(values, suffix=""):
    arr = sorted(set(values or []))
    if not arr:
        return ""
    parts = []
    start = prev = arr[0]
    for cur in arr[1:] + [None]:
        if cur == prev + 1:
            prev = cur
            continue
        parts.append(str(start) if start == prev else f"{start}-{prev}")
        start = prev = cur
    return ",".join(parts) + suffix


def lines_of(chunk):
    return [ln.strip() for ln in str(chunk or "").replace("\r", "").split("\n") if ln.strip()]


def clean_room(line):
    l = str(line or "").strip()
    m = re.match(r"^(.*?)[（(]([^（）()]*)[)）]\s*$", l)
    if m and m.group(2) and m.group(1).strip().startswith(m.group(2)):
        l = m.group(1).strip()
    return l


def parse_block(chunk):
    lines = lines_of(chunk)
    if not lines:
        return None
    c = {"name": "", "nature": "", "exam": "", "teacher": "", "room": "",
         "weeks": [], "periods": [], "note": "", "noTime": False}

    # 单行自由文本（如网课行）：拆出 课程名 / 教师 / 周次
    if len(lines) == 1:
        line = lines[0]
        wk = re.search(r"\[([^\]]*周[^\]]*)\]\s*$", line)
        if wk:
            c["weeks"] = parse_ranges(wk.group(1))
            line = line[: wk.start()].strip()
        tc = re.search(r"(\S+-\S+\[\S*\]\s*;?)\s*$", line)
        if tc:
            c["teacher"] = re.sub(r"[\[\];\s]+$", "", tc.group(1)).split("-")[0].strip()
            line = line[: tc.start()].strip()
        lines[0] = line
        if not line:
            return None

    name_line = lines[0]

    def nature_repl(m):
        c["nature"] = m.group(1)
        return ""

    def exam_repl(m):
        c["exam"] = m.group(1)
        return ""

    name_line = re.sub(r"[（(](必修|选修|限选|任选|公选|专选|实践|实训)[)）]", nature_repl, name_line)
    name_line = re.sub(r"[\[【](考试|考查)[\]】]", exam_repl, name_line)
    c["name"] = re.sub(r"\s+", " ", name_line).strip()

    for line in lines[1:]:
        wp = re.match(r"^\s*\[([^\]]*)\]\s*(?:\[([^\]]*)\])?\s*$", line)
        if wp and (("周" in wp.group(1) or "週" in wp.group(1)) or re.match(r"^\d", wp.group(1))):
            c["weeks"] = parse_ranges(wp.group(1))
            if wp.group(2):
                c["periods"] = parse_ranges(wp.group(2))
            continue
        w_only = re.match(r"^\s*\[([^\]]*周[^\]]*)\]\s*$", line)
        if w_only:
            c["weeks"] = parse_ranges(w_only.group(1))
            continue
        teacher_match = re.match(r"^([^[\]\-–—]+?)(?:\s*[-–—]\s*\S+)?(\[[^\]]*\])?;?\s*$", line)
        if teacher_match and ("[" in line or re.search(r"主讲|教师|老师|辅导", line)):
            c["teacher"] = teacher_match.group(1).strip()
            continue
        if re.match(r"^(组班|班级|合班)[：:]", line):
            c["note"] = re.sub(r"\s+", " ", line).strip()
            continue
        if re.match(r"^(学生|人数)[：:]", line):
            continue
        if re.match(r"^(教室|地点|上课地点|周次|节次)[：:]", line):
            c["room"] = re.sub(r"^(教室|地点|上课地点)[：:]\s*", "", line)
            continue
        if not c["room"] and re.search(r"楼|教室|实验室|机房|馆|场|-\d|\d{3,}", line) and not re.match(r"^\s*\[", line):
            c["room"] = clean_room(line)
            continue
        if not c["teacher"] and re.search(r"老师|教授|讲师|主讲", line):
            c["teacher"] = re.sub(r"老师|教授|讲师|主讲", "", line)
            c["teacher"] = re.sub(r"[\[\];：:]+", "", c["teacher"]).strip()
            continue
        if not c["room"]:
            c["room"] = clean_room(line)
    if not c["name"]:
        return None
    if not c["periods"]:
        c["noTime"] = True
    return c


def parse_blocks_text(text, day):
    out = []
    for chunk in re.split(r"学生[：:]\s*\d+\s*人", str(text or "")):
        c = parse_block(chunk)
        if not c:
            continue
        c["day"] = day or 0
        if c["noTime"] and day:
            c["noTime"] = False
        out.append(c)
    return out


def period_from_label(label):
    s = str(label or "").strip()
    if not s:
        return None
    pair = re.search(r"[第]?\s*(\d+)\s*[-–—~～]\s*(\d+)\s*节?", s)
    if pair:
        return (int(pair.group(1)), int(pair.group(2)))
    single = re.search(r"(\d+)\s*节?", s)
    if single and "周" not in s:
        return (int(single.group(1)), int(single.group(1)))
    if re.fullmatch(r"[一二三四五六七八九十]", s):
        n = NUM_CHAR[s]
        return (n * 2 - 1, n * 2)
    return None


def dedup(courses):
    merged = {}
    order = []
    for raw in courses:
        key = (raw.get("day", 0), raw["name"], raw.get("nature", ""), raw.get("teacher", ""),
               raw.get("room", ""), raw.get("note", ""), 1 if raw.get("noTime") else 0)
        if key not in merged:
            item = {
                "day": raw.get("day", 0),
                "name": raw["name"],
                "nature": raw.get("nature", ""),
                "exam": raw.get("exam", ""),
                "teacher": raw.get("teacher", ""),
                "room": raw.get("room", ""),
                "note": raw.get("note", ""),
                "noTime": bool(raw.get("noTime")),
                "weeks": set(raw.get("weeks") or []),
                "periods": set(raw.get("periods") or []),
            }
            merged[key] = item
            order.append(key)
        else:
            merged[key]["weeks"] |= set(raw.get("weeks") or [])
            merged[key]["periods"] |= set(raw.get("periods") or [])
    result = []
    for key in order:
        c = merged[key]
        weeks = sorted(c["weeks"])
        periods = sorted(c["periods"])
        result.append({
            "day": c["day"], "name": c["name"], "nature": c["nature"], "exam": c["exam"],
            "teacher": c["teacher"], "room": c["room"], "note": c["note"], "noTime": c["noTime"],
            "weeks": weeks, "weekText": summarize(weeks, "周"),
            "periods": periods, "periodText": summarize(periods, "节"),
        })
    result.sort(key=lambda x: (x["day"], x["noTime"], min(x["periods"]) if x["periods"] else 0, x["name"]))
    return result


def parse_sheet(ws):
    grid = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
    header_row = -1
    day_map = {}
    for r in range(min(len(grid), 8)):
        row = grid[r]
        found = {}
        for c, v in enumerate(row):
            d = match_day(v)
            if d:
                found[c] = d
        if len(found) >= 2:
            header_row = r
            day_map = found
            break
    if header_row < 0:
        raise ValueError("没有找到星期表头，无法识别课表")

    courses = []
    max_period = 0
    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        label = str(row[0] or "").strip() if row else ""
        pair = period_from_label(label)
        is_other = bool(re.search(r"占周|不占时间|网课|慕课|线上|其他|备注", label))
        for c, day in day_map.items():
            if c >= len(row):
                continue
            text = str(row[c] or "").strip()
            if not text:
                continue
            if is_other or re.search(r"占周不占时间|线上课程|网络课", text):
                for x in parse_blocks_text(text, 0):
                    x["noTime"] = True
                    x["day"] = 0
                    courses.append(x)
                continue
            parsed = parse_blocks_text(text, day) if "学生：" in text else None
            if not parsed:
                c0 = {
                    "name": lines_of(text)[0].replace("\n", " ") if lines_of(text) else "",
                    "nature": "", "exam": "", "teacher": "", "room": "",
                    "weeks": [], "periods": [], "note": "", "noTime": False,
                    "day": day,
                }
                for ln in lines_of(text)[1:]:
                    if re.search(r"老师|教授|讲师|主讲", ln):
                        c0["teacher"] = re.sub(r"老师|教授|讲师|主讲", "", ln).strip()
                    elif re.search(r"楼|教室|实验室|机房|馆|场|\d{3,}", ln):
                        c0["room"] = clean_room(ln)
                    else:
                        c0["note"] = ln
                if pair:
                    c0["periods"] = list(range(pair[0], pair[1] + 1))
                else:
                    c0["noTime"] = True
                parsed = [c0]
            for x in parsed:
                if x["periods"]:
                    max_period = max(max_period, max(x["periods"]))
            courses.extend(parsed)
        if pair and pair[1] > max_period:
            max_period = pair[1]
    return dedup(courses), max_period or 10


def main():
    ap = argparse.ArgumentParser(description="教务系统课表 Excel -> data.js")
    ap.add_argument("input", help="输入的 .xlsx 文件")
    ap.add_argument("-o", "--output", default="data.js", help="输出的 data.js 路径")
    ap.add_argument("--semester-start", default="2026-09-07", help="第1周的周一日期 YYYY-MM-DD")
    ap.add_argument("--days", type=int, default=7, help="每周显示天数 5/6/7")
    ap.add_argument("--json", help="同时输出一份可读的 JSON 用于检查")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.input, data_only=True)
    ws = wb.worksheets[0]
    courses, max_period = parse_sheet(ws)
    if not courses:
        print("没有解析到任何课程，请检查文件格式")
        sys.exit(1)
    period_count = max(12, ((max_period + 1) // 2) * 2)
    data = {
        "settings": {
            "semesterStart": args.semester_start,
            "daysPerWeek": args.days,
            "periodCount": period_count,
        },
        "courses": courses,
    }
    js = "// 由 convert_excel.py 自动生成，请勿手工修改\nwindow.BUILTIN_SCHEDULE = " + json.dumps(
        data, ensure_ascii=False, indent=2) + ";\n"
    Path(args.output).write_text(js, encoding="utf-8")
    print(f"OK: parsed {len(courses)} course entries -> {args.output}")
    if args.json:
        Path(args.json).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"OK: readable JSON -> {args.json}")


if __name__ == "__main__":
    main()
